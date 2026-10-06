"""「내집마련.xlsx」 → 공급유형 자격·배점 사양. 순수 파서 — 네트워크·DB 없음.

사용자가 만든 자격진단 엑셀 한 벌이 원본이다. 여기서 읽는 시트는 셋.

    신청자격DB   32개 공급유형의 나이·혼인·계층·무주택·소득·자산·자동차·지역 기준
    순위가점DB   같은 유형의 선정방식과 배점 항목. 배점 칸은 "24↑3/12↑2/6↑1" 같은 사람 말이다
    소득기준     도시근로자 가구원수별 월평균소득과, 서울/연접지역 시군구 목록

엑셀은 "제한 없음"을 999(소득%)·999999(자산·자동차만원)로 적는다. 여기서 None으로 눕힌다 —
0(자동차를 아예 못 가짐)과 무제한을 같은 값으로 두면 진단이 뒤집힌다.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl

SHEET_RULE = "신청자격DB"
SHEET_SCORE = "순위가점DB"
SHEET_INCOME = "소득기준"

NO_LIMIT_PCT = 999
NO_LIMIT_MAN = 999_999

# 소득기준 시트에서 지역 목록이 있는 칸 (1-based). 헤더 두 줄은 라벨과 예시라 건너뛴다
INCOME_HEADER_ROW = 3
COL_SEOUL, COL_NEAR = 11, 12

# 카테고리(+유형명) → housing_type 이넘. 못 잇는 유형은 None으로 둔다
_HOUSING_TYPE = {
    "행복주택": "행복주택",
    "청년안심주택(공공)": "청년안심주택",
    "청년안심주택(민간)": "청년안심주택",
    "매입임대": "매입임대",
    "전세임대": "전세임대",
    "장기전세": "장기전세",
    "든든전세": "든든전세",
    "전세형매입임대": "매입임대",
    "공공지원민간임대": "공공지원민간임대",
    "수요자맞춤형": None,
}
_HOUSING_TYPE_BY_NAME = {"국민임대": "국민임대", "영구임대": "영구임대", "재개발임대": "재개발임대"}

# 배점 시트에서 typed 컬럼으로 따로 빼는 칸. 나머지는 score jsonb로 통째 들어간다
_SCORE_TYPED = {"ID", "카테고리", "유형명", "선정방식",
                "우선1순위조건", "우선2순위조건", "우선3순위조건", "우선4순위조건",
                "일반1순위조건", "일반2순위조건", "일반3순위조건"}


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _int(value: Any, default: int = 0) -> int:
    s = _text(value)
    return int(float(s)) if s else default


def _limit(value: Any, sentinel: int) -> int | None:
    """제한 없음 표시(999·999999)를 None으로. 0은 진짜 0이라 그대로 둔다."""
    s = _text(value)
    if not s:
        return None
    n = int(float(s))
    return None if n >= sentinel else n


def _list(value: Any) -> list[str]:
    """`|`로 늘어놓은 칸을 배열로."""
    return [p.strip() for p in _text(value).split("|") if p.strip()]


def _rows(ws) -> list[dict[str, Any]]:
    """첫 줄을 헤더로 삼아 dict 목록으로. 빈 헤더 칸은 버린다."""
    it = ws.iter_rows(values_only=True)
    header = [_text(h) for h in next(it)]
    out = []
    for row in it:
        if not row or not _text(row[0]):
            continue
        out.append({h: row[i] for i, h in enumerate(header) if h and i < len(row)})
    return out


def housing_type_of(category: str, name: str) -> str | None:
    if category in _HOUSING_TYPE:
        return _HOUSING_TYPE[category]
    return _HOUSING_TYPE_BY_NAME.get(name)


# 비고 칸은 엑셀 작성자의 메모라 「2순위:본인+부모합산100%+자산34500」 「자동차 심사X」 꼴이다.
# 공고 상세 카드에 그대로 나가 읽기 힘들었다(사용자 지적 2026-10-06) — 원문 그대로를 키로 문장을 둔다.
# 신혼 메모의 70%/130%는 외벌이 기준이다(docs/eligibility-audit.md, 2026-10-06 공고문 대조).
# 엑셀 메모가 바뀌면 키가 안 맞아 원문이 그대로 나간다(뜻을 지어내지 않는다). 그때 여기에 한 줄 더한다.
# 화면 문자열이라 가운뎃점을 쓰지 않는다. 줄바꿈(\n)은 화면이 그대로 살린다.
NOTE_TEXT = {
    "미성년도 신청가능. 본인+부모소득합산": "미성년자도 신청할 수 있다. 소득은 본인과 부모를 합산해 본다",
    "미혼만 신청가능, 본인만 심사": "미혼만 신청할 수 있고, 소득과 자산은 본인 것만 본다",
    "맞벌이120%. 예비신혼/한부모 포함": "맞벌이 가구는 소득 120% 이하. 예비신혼부부와 한부모가족도 신청할 수 있다",
    "증명서 발급 필수": "주거급여 수급자 증명서를 내야 한다",
    "2순위:본인+부모합산100%+자산34500 \n3순위:본인100%+25100":
        "2순위: 본인과 부모 소득 합산 100% 이하, 총자산 3억 4,500만 원 이하\n"
        "3순위: 본인 소득 100% 이하, 총자산 2억 5,100만 원 이하",
    "맞벌이90%. 소득70%이하": "외벌이 가구는 소득 70% 이하, 맞벌이 가구는 90% 이하",
    "맞벌이200%. 소득130%이하": "외벌이 가구는 소득 130% 이하, 맞벌이 가구는 200% 이하",
    "맞벌이200%. 소득130%이하, 자동차 심사X":
        "외벌이 가구는 소득 130% 이하, 맞벌이 가구는 200% 이하. 자동차는 심사하지 않는다",
    "부부합산 소득 120%": "부부 합산 소득 120% 이하",
    "소득/자산/지역 요건 없이 추첨. 나이와 혼인기간, 무주택, 자동차는 청년안심주택 공통 요건으로 본다":
        "소득과 자산, 지역 요건 없이 추첨한다. 나이와 혼인기간, 무주택, 자동차는 청년안심주택 공통 요건으로 본다",
    "소득/자산 무관, 추첨": "소득과 자산을 보지 않고 추첨한다",
    "수급자/한부모/고령자/장애인": "수급자, 한부모가족, 고령자, 장애인 가구가 대상이다",
    "소득70%이하": "소득 70% 이하",
    "수급자/한부모/장애인 등": "수급자, 한부모가족, 장애인 가구 등이 대상이다",
    "수급자/한부모/차상위": "수급자, 한부모가족, 차상위계층이 먼저 뽑힌다",
    "맞벌이90%. 순위별선착순": "맞벌이 가구는 소득 90% 이하. 같은 순위 안에서는 선착순",
    "맞벌이200%. 자동차 심사X": "맞벌이 가구는 소득 200% 이하. 자동차는 심사하지 않는다",
    "미성년 자녀 2명 이상": "미성년 자녀가 2명 이상이어야 한다",
    "미리내집. 서울거주필수": "미리내집. 서울 거주자만 신청할 수 있다",
    "배점합산": "배점을 합산해 순위를 정한다",
    "자치구별상이": "자치구마다 조건이 다르다",
    "소득만 확인, 자산 및 자동차 심사X": "소득만 보고, 자산과 자동차는 심사하지 않는다",
}


def note_text(raw: str) -> str:
    return NOTE_TEXT.get(raw, raw)


def parse_supply_types(wb) -> list[dict[str, Any]]:
    """신청자격DB + 순위가점DB를 코드로 합쳐 supply_type 행으로."""
    scores = {_text(r["ID"]): r for r in _rows(wb[SHEET_SCORE])}
    out: list[dict[str, Any]] = []
    for order, r in enumerate(_rows(wb[SHEET_RULE]), start=1):
        code = _text(r["ID"])
        category, name = _text(r["카테고리"]), _text(r["유형명"])
        s = scores.get(code, {})
        out.append({
            "code": code,
            "category": category,
            "name": name,
            "housing_type": housing_type_of(category, name),
            "sort_order": order,
            "age_min": _int(r["최소나이"]),
            "age_max": _int(r["최대나이"], 999),
            "age_exempt": _list(r["나이면제조건"]),
            "marital": _text(r["혼인조건"]) or "무관",
            "marital_max_yr": _int(r["혼인MAX년"]),
            "newborn_exempt": _text(r["신생아혼인면제"]).upper() == "Y",
            "required_class": _list(r["필수계층"]),
            "homeless_scope": _text(r["무주택요건"]) or "세대원",
            "income_scope": _text(r["소득기준방식"]) or "무관",
            "income_pct": _limit(r["소득기준%"], NO_LIMIT_PCT),
            "asset_scope": _text(r["자산기준방식"]) or "무관",
            "asset_limit_man": _limit(r["자산만원"], NO_LIMIT_MAN),
            "car_limit_man": _limit(r["자동차만원"], NO_LIMIT_MAN),
            "region_limit": _text(r["지역제한"]) or "전국",
            "birth_bonus": _text(r["출산자녀가산"]).upper() == "Y",
            "note": note_text(_text(r["비고"])) or None,
            "ranking_method": _text(s.get("선정방식")) or None,
            "ranks": [x for x in (_text(s.get(f"우선{i}순위조건")) for i in (1, 2, 3, 4)) if x],
            "general_ranks": [x for x in (_text(s.get(f"일반{i}순위조건")) for i in (1, 2, 3)) if x],
            "score": {k: _text(v) for k, v in s.items() if k not in _SCORE_TYPED and _text(v)},
        })
    return out


def parse_income_standard(wb, *, year: int) -> list[dict[str, Any]]:
    """소득기준 시트의 가구원수 × % 표. 헤더 줄에서 % 칸 위치를 읽는다."""
    ws = wb[SHEET_INCOME]
    header = [_text(c) for c in next(ws.iter_rows(min_row=INCOME_HEADER_ROW, max_row=INCOME_HEADER_ROW, values_only=True))]
    pct_cols = {i: int(h.rstrip("%")) for i, h in enumerate(header) if h.endswith("%") and h[:-1].isdigit()}
    out = []
    for row in ws.iter_rows(min_row=INCOME_HEADER_ROW + 1, values_only=True):
        size = _text(row[0])
        if not size.isdigit():
            continue
        for i, pct in pct_cols.items():
            won = _text(row[i])
            if won:
                out.append({"year": year, "household": int(size), "pct": pct, "monthly_won": int(float(won))})
    return out


def parse_region_tiers(wb) -> list[dict[str, Any]]:
    """소득기준 시트 오른쪽에 붙어 있는 서울/연접지역 목록. 인천광역시처럼 시도 단위가 섞여 있다."""
    ws = wb[SHEET_INCOME]
    out: list[dict[str, Any]] = []
    for col, tier in ((COL_SEOUL, "서울"), (COL_NEAR, "연접")):
        for (cell,) in ws.iter_rows(min_row=INCOME_HEADER_ROW - 1, min_col=col, max_col=col, values_only=True):
            name = _text(cell)
            if not name or name in ("서울특별시", "연접지역", "기타지역", "기타"):
                continue
            out.append({"name": name, "kind": "sido" if name.endswith(("광역시", "특별시", "도")) else "sigungu", "tier": tier})
    return out


def parse_workbook(path: Path, *, year: int) -> dict[str, Any]:
    # read_only 모드는 같은 시트를 두 번 훑을 때 커서가 꼬인다. 파일이 작아 통째로 연다
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        return {
            "source": path.name,
            "income_year": year,
            "supply_types": parse_supply_types(wb),
            "income_standard": parse_income_standard(wb, year=year),
            "region_tiers": parse_region_tiers(wb),
        }
    finally:
        wb.close()
